"""
test_3c_ic_excel.py
-------------------
Generate an Excel workbook that manually verifies CrossSectionIC.

This uses the same factor and IC parameters as test_3b:
  factor: ColumnRef(DataColumn.OPEN_ADJUSTED)
  $F:     1min
  RE:     NextReturns(SC=OPEN_ADJUSTED, RF=2min, S=0)
  Lag:    0

Workbook layout:
  - one sheet per product: signal_time, backend FE, backend RE, data_present
  - IC: Excel formulas calculate cross-sectional Spearman IC row-by-row
  - BACKEND_COMPARE: Excel formulas compare backend IC vs Excel-calculated IC
  - IC_STATS: Excel formulas compare ic_stats for backend IC vs Excel IC
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import sys
from pathlib import Path
from typing import Iterable

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from sources.LocalCNFutures.CNFutures import CNFutures
from tools.data.types.DataColumn import DataColumn
from tools.factors.FactorExpr import ColumnRef
from tools.factors.FactorFamily import FactorFamily
from tools.factors.FactorTester import FactorTester
from tools.factors.tests.NextReturns import NextReturns
from tools.factors.tests.single_factor_test.ic import run_ic_for_factor

from tests.calc import (
    HEADER_FILL,
    MAIN_MINK_DIR,
    REMARK_FILL,
    REMARK_FONT,
    TEST_3B_DIR,
    remark_height,
)


DEFAULT_PRODUCTS = ['A', 'B', 'CU', 'AU', 'RB', 'I', 'J', 'JM', 'PP', 'TA']
DEFAULT_START_DATE = '2025-01-02'
DEFAULT_DAYS = 30
DEFAULT_RF_MINUTES = 2
TOLERANCE = 1e-10

OUT_DIR = TEST_3B_DIR.parent / 'test_3c'
OUT_PATH = OUT_DIR / 'ic_excel_check.xlsx'

DATE_TIME_FORMAT = 'yyyy-mm-dd hh:mm:ss'
NUMBER_FORMAT = '0.0000000000'
IC_SHEET = 'IC'
COMPARE_SHEET = 'BACKEND_COMPARE'
STATS_SHEET = 'IC_STATS'
IC_DATA_START_ROW = 4
COMPARE_DATA_START_ROW = 10


class _OpenAdjustedFactor(FactorFamily):
    source_freq = '1m'

    @staticmethod
    def factor_expr():
        return ColumnRef(DataColumn.OPEN_ADJUSTED)


def _product_name(product_id: str) -> str:
    if '.' in product_id:
        return product_id
    candidates = [
        path.stem
        for path in MAIN_MINK_DIR.glob(f'{product_id}.*.parquet')
        if path.name.split('.', 1)[0] == product_id
    ]
    if not candidates:
        raise RuntimeError(f'Cannot infer exchange for product {product_id}; pass full name like A.DCE')
    if len(candidates) > 1:
        raise RuntimeError(f'Multiple products match {product_id}: {", ".join(sorted(candidates))}')
    return candidates[0]


def _normalise_ts(value) -> pd.Timestamp:
    ts = pd.Timestamp(value)
    if ts.tzinfo is not None:
        ts = ts.tz_convert('Asia/Shanghai').tz_localize(None)
    return ts


def _signal_time(idx) -> pd.Timestamp:
    return _normalise_ts(idx[-1] if isinstance(idx, tuple) else idx)


def _normalise_number(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass
    value = float(value)
    if math.isnan(value) or math.isinf(value):
        return None
    return value


def _normalise_stat_value(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass
    value = float(value)
    if math.isnan(value):
        return None
    if math.isinf(value):
        return 'inf'
    return value


def _safe_sheet_name(name: str, used: set[str]) -> str:
    base = name.replace('.', '_').replace('/', '_')[:31] or 'PRODUCT'
    candidate = base
    suffix = 1
    while candidate in used:
        tail = f'_{suffix}'
        candidate = f'{base[:31 - len(tail)]}{tail}'
        suffix += 1
    used.add(candidate)
    return candidate


def _write_remark(ws, text: str, max_col: int) -> None:
    ws.cell(1, 1, text)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.row_dimensions[1].height = remark_height(text)
    cell = ws.cell(1, 1)
    cell.fill = REMARK_FILL
    cell.font = REMARK_FONT
    cell.alignment = Alignment(wrap_text=True, vertical='top')


def _write_header(ws, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row, col, header)
        cell.fill = HEADER_FILL
        cell.font = REMARK_FONT


def _resolve_series(table: pd.DataFrame, product) -> pd.Series:
    if product in table.columns:
        return table[product]
    for col in table.columns:
        if getattr(col, 'name', None) == product.name or str(col) == product.name:
            return table[col]
    raise RuntimeError(f'{product.name}: product column not found')


def _build_backend_tables(products: list[CNFutures], start: pd.Timestamp, end: pd.Timestamp, rf_minutes: int):
    tester = FactorTester(products=products, time_range=(start, end), logger_file=False)
    factor = _OpenAdjustedFactor().get_factor(**{'$F': '1min', '$Rev': '0'})
    factor.evaluate(products)
    next_returns = NextReturns().get_factor(
        SC=DataColumn.OPEN_ADJUSTED,
        RF=f'{rf_minutes}min',
        S=0,
        **{'$F': '1min', '$Rev': '0'},
    )
    params = {
        'FE': factor,
        'RE': next_returns,
        'Lag': 0,
        '$F': '1min',
    }
    _factors, ic_series, stats, re_table, fe_table, mask = run_ic_for_factor(tester, params, [factor])

    ic_series = pd.Series(pd.to_numeric(ic_series, errors='coerce'))
    signal_times = pd.DatetimeIndex([_signal_time(idx) for idx in ic_series.index])

    product_data = {}
    for product in products:
        fe = _resolve_series(fe_table, product).copy()
        re = _resolve_series(re_table, product).copy()
        present = _resolve_series(mask, product).copy() if not mask.empty else pd.Series(True, index=fe.index)

        fe.index = pd.DatetimeIndex([_signal_time(idx) for idx in fe.index])
        re.index = pd.DatetimeIndex([_signal_time(idx) for idx in re.index])
        present.index = pd.DatetimeIndex([_signal_time(idx) for idx in present.index])

        product_data[product.name] = pd.DataFrame({
            'signal_time': signal_times,
            'fe': fe.reindex(signal_times).to_numpy(),
            're': re.reindex(signal_times).to_numpy(),
            'data_present': present.reindex(signal_times).fillna(False).astype(bool).to_numpy(),
        })

    backend_ic = pd.DataFrame({
        'signal_time': signal_times,
        'backend_ic': ic_series.to_numpy(),
    })
    return backend_ic, product_data, stats


def _write_product_sheet(ws, product_name: str, data: pd.DataFrame) -> None:
    _write_remark(
        ws,
        f'{product_name}: backend FE/RE intermediate values used by Excel IC formulas.',
        4,
    )
    _write_header(ws, 2, ['signal_time', 'FE_OPEN_ADJUSTED', 'RE_OPEN_TO_OPEN', 'data_present'])
    ws.freeze_panes = 'A3'
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 14

    for row_idx, row in enumerate(data.itertuples(index=False), start=3):
        ws.cell(row_idx, 1, row.signal_time.to_pydatetime()).number_format = DATE_TIME_FORMAT
        ws.cell(row_idx, 2, _normalise_number(row.fe)).number_format = NUMBER_FORMAT
        ws.cell(row_idx, 3, _normalise_number(row.re)).number_format = NUMBER_FORMAT
        ws.cell(row_idx, 4, bool(row.data_present))


def _write_ic_sheet(ws, backend_ic: pd.DataFrame, sheet_map: dict[str, str]) -> tuple[int, int]:
    products = list(sheet_map)
    fe_start_col = 3
    re_start_col = fe_start_col + len(products)
    fe_rank_start_col = re_start_col + len(products)
    re_rank_start_col = fe_rank_start_col + len(products)
    valid_col = re_rank_start_col + len(products)
    excel_ic_col = valid_col + 1
    max_col = excel_ic_col

    _write_remark(
        ws,
        'Excel formula IC: each row calculates cross-sectional Spearman rank correlation from product FE/RE sheets.',
        max_col,
    )
    ws.cell(2, 1, 'tolerance')
    ws.cell(2, 2, TOLERANCE)
    ws.cell(2, 4, 'formula')
    ws.cell(2, 5, 'CORREL(FE_RANK, RE_RANK) over valid product columns')

    headers = ['signal_time', 'backend_ic']
    headers += [f'FE_{p}' for p in products]
    headers += [f'RE_{p}' for p in products]
    headers += [f'FE_RANK_{p}' for p in products]
    headers += [f'RE_RANK_{p}' for p in products]
    headers += ['valid_n', 'excel_ic']
    _write_header(ws, 3, headers)
    ws.freeze_panes = 'C4'
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 16

    fe_end_col = fe_start_col + len(products) - 1
    re_end_col = re_start_col + len(products) - 1
    fe_rank_end_col = fe_rank_start_col + len(products) - 1
    re_rank_end_col = re_rank_start_col + len(products) - 1

    for row_idx, row in enumerate(backend_ic.itertuples(index=False), start=IC_DATA_START_ROW):
        ws.cell(row_idx, 1, row.signal_time.to_pydatetime()).number_format = DATE_TIME_FORMAT
        ws.cell(row_idx, 2, _normalise_number(row.backend_ic)).number_format = NUMBER_FORMAT

        for offset, product in enumerate(products):
            col = fe_start_col + offset
            sheet = sheet_map[product]
            product_row = row_idx - 1
            ws.cell(
                row_idx, col,
                f'=IF(AND({sheet}!$D{product_row},ISNUMBER({sheet}!$B{product_row}),ISNUMBER({sheet}!$C{product_row})),'
                f'{sheet}!$B{product_row},"")'
            )
            ws.cell(row_idx, col).number_format = NUMBER_FORMAT
        for offset, product in enumerate(products):
            col = re_start_col + offset
            sheet = sheet_map[product]
            product_row = row_idx - 1
            ws.cell(
                row_idx, col,
                f'=IF(AND({sheet}!$D{product_row},ISNUMBER({sheet}!$B{product_row}),ISNUMBER({sheet}!$C{product_row})),'
                f'{sheet}!$C{product_row},"")'
            )
            ws.cell(row_idx, col).number_format = NUMBER_FORMAT

        fe_range = f'{ws.cell(row_idx, fe_start_col).coordinate}:{ws.cell(row_idx, fe_end_col).coordinate}'
        re_range = f'{ws.cell(row_idx, re_start_col).coordinate}:{ws.cell(row_idx, re_end_col).coordinate}'
        fe_rank_range = f'{ws.cell(row_idx, fe_rank_start_col).coordinate}:{ws.cell(row_idx, fe_rank_end_col).coordinate}'
        re_rank_range = f'{ws.cell(row_idx, re_rank_start_col).coordinate}:{ws.cell(row_idx, re_rank_end_col).coordinate}'

        for offset, _product in enumerate(products):
            col = fe_rank_start_col + offset
            fe_cell = ws.cell(row_idx, fe_start_col + offset).coordinate
            re_cell = ws.cell(row_idx, re_start_col + offset).coordinate
            ws.cell(
                row_idx, col,
                f'=IF(AND(ISNUMBER({fe_cell}),ISNUMBER({re_cell})),'
                f'1+SUMPRODUCT(--ISNUMBER({fe_range}),--ISNUMBER({re_range}),--({fe_range}<{fe_cell}))'
                f'+(SUMPRODUCT(--ISNUMBER({fe_range}),--ISNUMBER({re_range}),--({fe_range}={fe_cell}))-1)/2,"")'
            )
            ws.cell(row_idx, col).number_format = NUMBER_FORMAT
        for offset, _product in enumerate(products):
            col = re_rank_start_col + offset
            fe_cell = ws.cell(row_idx, fe_start_col + offset).coordinate
            re_cell = ws.cell(row_idx, re_start_col + offset).coordinate
            ws.cell(
                row_idx, col,
                f'=IF(AND(ISNUMBER({fe_cell}),ISNUMBER({re_cell})),'
                f'1+SUMPRODUCT(--ISNUMBER({fe_range}),--ISNUMBER({re_range}),--({re_range}<{re_cell}))'
                f'+(SUMPRODUCT(--ISNUMBER({fe_range}),--ISNUMBER({re_range}),--({re_range}={re_cell}))-1)/2,"")'
            )
            ws.cell(row_idx, col).number_format = NUMBER_FORMAT

        valid_formula = f'=SUMPRODUCT(--ISNUMBER({fe_range}),--ISNUMBER({re_range}))'
        ic_formula = f'=IF({ws.cell(row_idx, valid_col).coordinate}<=1,"",IFERROR(CORREL({fe_rank_range},{re_rank_range}),""))'
        ws.cell(row_idx, valid_col, valid_formula)
        ws.cell(row_idx, excel_ic_col, ic_formula).number_format = NUMBER_FORMAT

    for col in range(3, max_col + 1):
        ws.column_dimensions[ws.cell(3, col).column_letter].width = 15
    return valid_col, excel_ic_col


def _write_compare_sheet(ws, backend_ic_rows: int, valid_col: int, excel_ic_col: int) -> None:
    _write_remark(
        ws,
        'Excel formula comparison between backend CrossSectionIC output and IC sheet Excel formulas.',
        6,
    )
    ws.cell(2, 1, 'tolerance')
    ws.cell(2, 2, TOLERANCE)
    ws.cell(3, 1, 'missing_excel_ic')
    ws.cell(3, 2, f'=COUNTIF(F{COMPARE_DATA_START_ROW}:F{COMPARE_DATA_START_ROW + backend_ic_rows - 1},"MISSING")')
    ws.cell(4, 1, 'diff_rows')
    ws.cell(4, 2, f'=COUNTIF(F{COMPARE_DATA_START_ROW}:F{COMPARE_DATA_START_ROW + backend_ic_rows - 1},"FAIL")')
    ws.cell(5, 1, 'max_abs_diff')
    ws.cell(5, 2, f'=MAX(D{COMPARE_DATA_START_ROW}:D{COMPARE_DATA_START_ROW + backend_ic_rows - 1})')
    ws.cell(6, 1, 'overall')
    ws.cell(6, 2, '=IF(AND(B3=0,B4=0),"PASS","FAIL")')

    _write_header(ws, COMPARE_DATA_START_ROW - 1, ['signal_time', 'backend_ic', 'excel_ic', 'abs_diff', 'valid_n', 'status'])
    ws.freeze_panes = f'A{COMPARE_DATA_START_ROW}'
    ws.column_dimensions['A'].width = 22
    for col in 'BCDEF':
        ws.column_dimensions[col].width = 16

    valid_letter = ws.cell(1, valid_col).column_letter
    excel_ic_letter = ws.cell(1, excel_ic_col).column_letter
    for out_row in range(COMPARE_DATA_START_ROW, COMPARE_DATA_START_ROW + backend_ic_rows):
        ic_row = IC_DATA_START_ROW + (out_row - COMPARE_DATA_START_ROW)
        ws.cell(out_row, 1, f'=IC!A{ic_row}').number_format = DATE_TIME_FORMAT
        ws.cell(out_row, 2, f'=IC!B{ic_row}').number_format = NUMBER_FORMAT
        ws.cell(out_row, 3, f'=IC!{excel_ic_letter}{ic_row}').number_format = NUMBER_FORMAT
        ws.cell(
            out_row, 4,
            f'=IF(AND(B{out_row}="",C{out_row}=""),0,IF(OR(B{out_row}="",C{out_row}=""),"",ABS(B{out_row}-C{out_row})))'
        ).number_format = NUMBER_FORMAT
        ws.cell(out_row, 5, f'=IC!{valid_letter}{ic_row}')
        ws.cell(
            out_row, 6,
            f'=IF(AND(B{out_row}="",C{out_row}=""),"PASS",'
            f'IF(C{out_row}="","MISSING",IF(B{out_row}="","FAIL",IF(D{out_row}<=$B$2,"PASS","FAIL"))))'
        )


def _write_stats_sheet(ws, backend_ic_rows: int, excel_ic_col: int, backend_stats: pd.Series) -> None:
    _write_remark(
        ws,
        'IC summary comparison: backend_stat is the backend ic_stats() output; excel_stat is calculated by Excel formulas.',
        6,
    )
    ws.cell(2, 1, 'tolerance')
    ws.cell(2, 2, TOLERANCE)
    ws.cell(3, 1, 'compare_overall')
    ws.cell(3, 2, '=BACKEND_COMPARE!B6')
    ws.cell(4, 1, 'stats_overall')
    ws.cell(4, 2, '=IF(COUNTIF(F8:F15,"FAIL")=0,"PASS","FAIL")')

    _write_header(ws, 7, ['metric', 'backend_stat', 'excel_stat', 'abs_diff', 'tolerance', 'status'])
    ws.freeze_panes = 'A8'
    for col in 'ABCDEF':
        ws.column_dimensions[col].width = 18

    last_ic_row = IC_DATA_START_ROW + backend_ic_rows - 1
    excel_ic_letter = ws.cell(1, excel_ic_col).column_letter
    backend_range = f'IC!$B${IC_DATA_START_ROW}:$B${last_ic_row}'
    excel_range = f'IC!${excel_ic_letter}${IC_DATA_START_ROW}:${excel_ic_letter}${last_ic_row}'

    def ac1_formula(value_range: str) -> str:
        return (
            f'=IFERROR(LET(vals,FILTER({value_range},ISNUMBER({value_range})),'
            'n,ROWS(vals),'
            'IF(n<=2,"",IFERROR(CORREL(DROP(vals,-1),DROP(vals,1)),""))),"")'
        )

    def half_life_formula(value_range: str) -> str:
        return (
            f'=IFERROR(LET(vals,FILTER({value_range},ISNUMBER({value_range})),'
            'n,ROWS(vals),'
            'IF(n<=2,"",'
            'LET(maxLag,MIN(20,MAX(1,QUOTIENT(n,2)-1)),'
            'lags,SEQUENCE(maxLag),'
            'acfs,MAP(lags,LAMBDA(k,IFERROR(CORREL(DROP(vals,-k),DROP(vals,k)),""))),'
            'pos,IFERROR(XMATCH(TRUE,IF(ISNUMBER(acfs),acfs<0.5,FALSE)),""),'
            'IF(pos="","inf",'
            'LET(curr,INDEX(acfs,pos),'
            'prev,IF(pos=1,1,INDEX(acfs,pos-1)),'
            'lag,INDEX(lags,pos),'
            'lag-1+IF(curr=prev,0,(0.5-prev)/(curr-prev)))))))),"")'
        )

    metrics = [
        ('mean', _normalise_stat_value(backend_stats.get('mean')), f'=IFERROR(AVERAGE({excel_range}),"")'),
        ('std', _normalise_stat_value(backend_stats.get('std')), f'=IFERROR(STDEV.S({excel_range}),"")'),
        ('IR', _normalise_stat_value(backend_stats.get('IR')), '=IFERROR(IF(C9=0,"",C8/C9),"")'),
        ('t_stat', _normalise_stat_value(backend_stats.get('t_stat')), f'=IFERROR(IF(OR(C9=0,COUNT({excel_range})<=1),"",C8/(C9/SQRT(COUNT({excel_range})))),"")'),
        ('max', _normalise_stat_value(backend_stats.get('max')), f'=IF(COUNT({excel_range})=0,"",MAX({excel_range}))'),
        ('min', _normalise_stat_value(backend_stats.get('min')), f'=IF(COUNT({excel_range})=0,"",MIN({excel_range}))'),
        ('ac1', _normalise_stat_value(backend_stats.get('ac1')), ac1_formula(excel_range)),
        ('half_life', _normalise_stat_value(backend_stats.get('half_life')), half_life_formula(excel_range)),
    ]

    for row_idx, (metric, backend_value, excel_formula) in enumerate(metrics, start=8):
        ws.cell(row_idx, 1, metric)
        ws.cell(row_idx, 2, backend_value).number_format = NUMBER_FORMAT
        ws.cell(row_idx, 3, excel_formula).number_format = NUMBER_FORMAT
        ws.cell(
            row_idx, 4,
            f'=IF(AND(B{row_idx}="",C{row_idx}=""),0,IF(OR(B{row_idx}="",C{row_idx}=""),"",'
            f'IF(OR(ISTEXT(B{row_idx}),ISTEXT(C{row_idx})),'
            f'IF(B{row_idx}=C{row_idx},0,""),ABS(B{row_idx}-C{row_idx}))))'
        ).number_format = NUMBER_FORMAT
        ws.cell(row_idx, 5, TOLERANCE).number_format = NUMBER_FORMAT
        ws.cell(
            row_idx, 6,
            f'=IF(AND(B{row_idx}="",C{row_idx}=""),"PASS",'
            f'IF(D{row_idx}="","FAIL",IF(D{row_idx}<=E{row_idx},"PASS","FAIL")))'
        )


def build_workbook(
    products: Iterable[str],
    start_date: str,
    days: int,
    rf_minutes: int,
    out_path: Path,
) -> Path:
    start = pd.Timestamp(start_date)
    end = start + pd.Timedelta(days=days)
    product_objs = [CNFutures(_product_name(product)) for product in products]

    backend_ic, product_data, backend_stats = _build_backend_tables(product_objs, start, end, rf_minutes)
    if backend_ic.empty:
        raise RuntimeError('No backend IC rows in selected date window')

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    active = wb.active
    if active is not None:
        wb.remove(active)

    used_sheets = set()
    sheet_map: dict[str, str] = {}
    for product in product_objs:
        sheet_name = _safe_sheet_name(product.code, used_sheets)
        sheet_map[product.name] = sheet_name
        ws = wb.create_sheet(sheet_name)
        _write_product_sheet(ws, product.name, product_data[product.name])

    ic_ws = wb.create_sheet(IC_SHEET)
    valid_col, excel_ic_col = _write_ic_sheet(ic_ws, backend_ic, sheet_map)

    compare_ws = wb.create_sheet(COMPARE_SHEET)
    _write_compare_sheet(compare_ws, len(backend_ic), valid_col, excel_ic_col)

    stats_ws = wb.create_sheet(STATS_SHEET)
    _write_stats_sheet(stats_ws, len(backend_ic), excel_ic_col, backend_stats)

    wb.save(out_path)
    wb.close()
    return out_path


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description='test_3c: Excel-formula IC verification workbook.')
    parser.add_argument('products', nargs='*', help='Optional product ids; default is about 10 liquid products.')
    parser.add_argument('--start-date', default=DEFAULT_START_DATE, help=f'Default: {DEFAULT_START_DATE}')
    parser.add_argument('--days', type=int, default=DEFAULT_DAYS, help=f'Default: {DEFAULT_DAYS}')
    parser.add_argument('--rf-minutes', type=int, default=DEFAULT_RF_MINUTES, help=f'Default: {DEFAULT_RF_MINUTES}')
    parser.add_argument('--output', type=Path, default=OUT_PATH, help=f'Default: {OUT_PATH}')
    args = parser.parse_args(argv)

    products = args.products or DEFAULT_PRODUCTS
    print('test_3c: Excel-formula IC verification')
    print(f'Products: {", ".join(products)}')
    print(f'Window  : {args.start_date} + {args.days} days')
    print(f'RF      : {args.rf_minutes}min')
    print(f'Output  : {args.output}')
    out = build_workbook(products, args.start_date, args.days, args.rf_minutes, args.output)
    print(f'OK: {out}')


def test_test3c_workbook_smoke():
    out_path = OUT_DIR / '_smoke_ic_excel_check.xlsx'
    build_workbook(['A', 'B'], DEFAULT_START_DATE, 3, DEFAULT_RF_MINUTES, out_path)
    wb = load_workbook(out_path, read_only=False, data_only=False)
    try:
        assert IC_SHEET in wb.sheetnames
        assert COMPARE_SHEET in wb.sheetnames
        assert STATS_SHEET in wb.sheetnames
        assert len(wb.sheetnames) == 5
        assert wb[IC_SHEET].cell(3, 1).value == 'signal_time'
        assert wb[IC_SHEET].cell(3, 2).value == 'backend_ic'
        assert 'ISNUMBER' in str(wb[IC_SHEET].cell(IC_DATA_START_ROW, 3).value)
        assert 'ISNUMBER' in str(wb[IC_SHEET].cell(IC_DATA_START_ROW, 5).value)
        assert wb[IC_SHEET].cell(3, 7).value == 'FE_RANK_A.DCE'
        rank_formula = str(wb[IC_SHEET].cell(IC_DATA_START_ROW, 7).value)
        assert 'SUMPRODUCT' in rank_formula
        assert '--ISNUMBER(C4:D4),--ISNUMBER(E4:F4)' in rank_formula
        excel_ic_formula = str(wb[IC_SHEET].cell(IC_DATA_START_ROW, 12).value)
        assert 'IFERROR(CORREL' in excel_ic_formula
        assert wb[COMPARE_SHEET].cell(COMPARE_DATA_START_ROW - 1, 6).value == 'status'
        assert 'AND(B10="",C10="")' in str(wb[COMPARE_SHEET].cell(COMPARE_DATA_START_ROW, 6).value)
        assert wb[STATS_SHEET].cell(8, 1).value == 'mean'
        assert wb[STATS_SHEET].cell(11, 1).value == 't_stat'
        assert isinstance(wb[STATS_SHEET].cell(8, 2).value, (int, float))
        assert str(wb[STATS_SHEET].cell(8, 3).value).startswith('=IFERROR(AVERAGE(IC!$')
        assert 'AND(B15="",C15="")' in str(wb[STATS_SHEET].cell(15, 6).value)
    finally:
        wb.close()


if __name__ == '__main__':
    main()
