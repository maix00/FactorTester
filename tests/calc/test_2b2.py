"""
test_2b2.py
-----------
Command-line adjustment_mul comparison for test_2a MAIN vs main_mink.

Unlike test_2b, this script does not create an Excel workbook and does not
require Excel/WPS automation.  It reads cached formula values from the
test_2a workbook, compares them with main_mink parquet data in Python, and
exits non-zero when any product differs.

Prerequisite:
  test_2a/{prod}.xlsx must already have been opened, calculated, and saved by
  WPS/Excel or by tests/calc/test_2c_wps_compute.py.
"""

from __future__ import annotations

import argparse
import os
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.utils import range_boundaries

from tests.calc import MAIN_MINK_DIR, TEST_2A_DIR


DEFAULT_WORKERS = max(1, min(4, os.cpu_count() or 1))
DEFAULT_TOLERANCE = 1e-8

MAIN_SHEET = 'MAIN'
MAIN_HEADER_ROW = 2
MAIN_DATA_START_ROW = 3

TRADE_TIME_COL = 'trade_time'
ADJ_MUL_COL = 'adjustment_mul'


@dataclass(frozen=True)
class ProductComparison:
    product: str
    test2a_rows: int
    main_mink_rows: int
    missing_trade_time: int
    diff_rows: int
    max_abs_diff: float | None
    duplicate_test2a_keys: int
    duplicate_mink_keys: int
    examples: list[str]

    @property
    def ok(self) -> bool:
        return (
            self.missing_trade_time == 0
            and self.diff_rows == 0
            and self.duplicate_test2a_keys == 0
            and self.duplicate_mink_keys == 0
        )


def _excel_ref_bounds(ws, cell_ref: str) -> tuple[int | None, int | None, int | None, int | None]:
    value = ws[cell_ref].value
    ref = getattr(value, 'ref', None)
    if ref:
        return range_boundaries(ref)  # type: ignore[return-value]
    return (ws[cell_ref].column, ws[cell_ref].row, ws.max_column, ws.max_row)


def _find_main_mink_file(prod: str) -> Path | None:
    candidates = [
        p for p in MAIN_MINK_DIR.glob(f'{prod}.*.parquet')
        if p.name.split('.', 1)[0] == prod
    ]
    if not candidates:
        return None
    if len(candidates) > 1:
        names = ', '.join(p.name for p in sorted(candidates))
        raise RuntimeError(f'multiple main_mink files found for {prod}: {names}')
    return candidates[0]


def _col_index(headers: list, name: str) -> int:
    try:
        return [str(h) if h is not None else '' for h in headers].index(name)
    except ValueError as exc:
        raise RuntimeError(f'MAIN missing required column: {name}') from exc


def _normalise_trade_time(value):
    if pd.isna(value):
        return None
    if isinstance(value, (int, float)):
        return from_excel(value)
    if hasattr(value, 'to_pydatetime'):
        return value.to_pydatetime()
    return value


def _time_key(series: pd.Series) -> pd.Series:
    normalised = series.map(_normalise_trade_time)
    return pd.to_datetime(normalised, errors='coerce').dt.strftime('%Y-%m-%d %H:%M:%S')


def _load_test2a_adj(src_path: Path) -> pd.DataFrame:
    formula_wb = load_workbook(src_path, read_only=True, data_only=False)
    value_wb = load_workbook(src_path, read_only=True, data_only=True)
    try:
        if MAIN_SHEET not in formula_wb.sheetnames or MAIN_SHEET not in value_wb.sheetnames:
            raise RuntimeError(f'{src_path.name} has no MAIN sheet')

        formula_ws = formula_wb[MAIN_SHEET]
        value_ws = value_wb[MAIN_SHEET]
        _min_col, _min_row, max_col, max_row = _excel_ref_bounds(formula_ws, 'A3')
        if max_col is None:
            raise RuntimeError(f'{src_path.name} MAIN sheet has no data')

        headers = [value_ws.cell(MAIN_HEADER_ROW, c).value for c in range(1, max_col + 1)]
        first_value = value_ws.cell(MAIN_DATA_START_ROW, 1).value
        if not headers or headers[0] is None or first_value is None:
            raise RuntimeError(
                f'{src_path.name} MAIN has no cached calculated values. '
                'Run test_2c_wps_compute.py first, then rerun test_2b2.'
            )

        trade_time_idx = _col_index(headers, TRADE_TIME_COL)
        adj_idx = _col_index(headers, ADJ_MUL_COL)

        rows: list[tuple] = []
        for row in value_ws.iter_rows(
            min_row=MAIN_DATA_START_ROW,
            max_row=max_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        ):
            trade_time = row[trade_time_idx]
            adj_mul = row[adj_idx]
            if trade_time is None:
                continue
            rows.append((trade_time, adj_mul))

        df = pd.DataFrame(rows, columns=[TRADE_TIME_COL, 'test2a_adjustment_mul'])
        df['_trade_time_key'] = _time_key(df[TRADE_TIME_COL])
        df['test2a_adjustment_mul'] = pd.to_numeric(df['test2a_adjustment_mul'], errors='coerce')
        return df
    finally:
        formula_wb.close()
        value_wb.close()


def _load_main_mink_adj(mink_path: Path) -> pd.DataFrame:
    df = pd.read_parquet(mink_path, columns=[TRADE_TIME_COL, ADJ_MUL_COL])
    out = df[[TRADE_TIME_COL, ADJ_MUL_COL]].copy()
    out = out.rename(columns={ADJ_MUL_COL: 'main_mink_adjustment_mul'})
    out['_trade_time_key'] = _time_key(out[TRADE_TIME_COL])
    out['main_mink_adjustment_mul'] = pd.to_numeric(
        out['main_mink_adjustment_mul'], errors='coerce'
    )
    return out


def _example_rows(compared: pd.DataFrame, limit: int) -> list[str]:
    if limit <= 0:
        return []

    bad = compared[
        compared['main_mink_adjustment_mul'].isna()
        | compared['abs_diff'].isna()
        | (compared['abs_diff'] > compared.attrs['tolerance'])
    ].head(limit)

    examples: list[str] = []
    for _idx, row in bad.iterrows():
        examples.append(
            f'{row["_trade_time_key"]}: test2a={row["test2a_adjustment_mul"]}, '
            f'main_mink={row["main_mink_adjustment_mul"]}, abs_diff={row["abs_diff"]}'
        )
    return examples


def compare_product(test2a_xlsx: Path, tolerance: float, example_limit: int) -> ProductComparison:
    prod = test2a_xlsx.stem
    if not test2a_xlsx.exists():
        raise RuntimeError(f'test_2a file not found: {test2a_xlsx}')

    mink_path = _find_main_mink_file(prod)
    if mink_path is None:
        raise RuntimeError(f'main_mink file not found for {prod}')

    test2a_df = _load_test2a_adj(test2a_xlsx)
    mink_df = _load_main_mink_adj(mink_path)

    duplicate_test2a = int(test2a_df['_trade_time_key'].duplicated().sum())
    duplicate_mink = int(mink_df['_trade_time_key'].duplicated().sum())
    if duplicate_test2a:
        test2a_df = test2a_df.drop_duplicates('_trade_time_key', keep='first')
    if duplicate_mink:
        mink_df = mink_df.drop_duplicates('_trade_time_key', keep='first')

    compared = test2a_df.merge(
        mink_df[['_trade_time_key', 'main_mink_adjustment_mul']],
        on='_trade_time_key',
        how='left',
    )
    compared['abs_diff'] = (
        compared['test2a_adjustment_mul'] - compared['main_mink_adjustment_mul']
    ).abs()
    compared.attrs['tolerance'] = tolerance

    missing = int(compared['main_mink_adjustment_mul'].isna().sum())
    diff_rows = int((compared['abs_diff'] > tolerance).fillna(False).sum())
    max_abs_diff_value = compared['abs_diff'].max(skipna=True)
    max_abs_diff = None if pd.isna(max_abs_diff_value) else float(max_abs_diff_value)

    return ProductComparison(
        product=prod,
        test2a_rows=len(test2a_df),
        main_mink_rows=len(mink_df),
        missing_trade_time=missing,
        diff_rows=diff_rows,
        max_abs_diff=max_abs_diff,
        duplicate_test2a_keys=duplicate_test2a,
        duplicate_mink_keys=duplicate_mink,
        examples=_example_rows(compared, example_limit),
    )


def _iter_test2a_files(products: Iterable[str] | None = None) -> list[Path]:
    if products:
        return [TEST_2A_DIR / f'{prod}.xlsx' for prod in products]
    return sorted(
        path for path in TEST_2A_DIR.glob('*.xlsx')
        if not path.name.startswith('~$') and not path.name.startswith('.')
    )


def _print_result(result: ProductComparison) -> None:
    status = 'OK' if result.ok else 'FAIL'
    max_diff = 'NA' if result.max_abs_diff is None else f'{result.max_abs_diff:.12g}'
    print(
        f'  {result.product}: {status} '
        f'(test2a_rows={result.test2a_rows}, main_mink_rows={result.main_mink_rows}, '
        f'missing={result.missing_trade_time}, diff={result.diff_rows}, '
        f'dup_test2a={result.duplicate_test2a_keys}, dup_mink={result.duplicate_mink_keys}, '
        f'max_abs_diff={max_diff})'
    )
    for example in result.examples:
        print(f'    - {example}')


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description='test_2b2: CLI comparison of test_2a MAIN adjustment_mul against main_mink.'
    )
    parser.add_argument('products', nargs='*', help='Optional product ids, e.g. A RB IF')
    parser.add_argument(
        '--workers',
        type=int,
        default=int(os.environ.get('TEST_2B2_WORKERS', DEFAULT_WORKERS)),
        help='Number of products to compare in parallel.',
    )
    parser.add_argument(
        '--tolerance',
        type=float,
        default=DEFAULT_TOLERANCE,
        help=f'Absolute tolerance for adjustment_mul comparison (default: {DEFAULT_TOLERANCE}).',
    )
    parser.add_argument(
        '--examples',
        type=int,
        default=5,
        help='Number of missing/different row examples to print per failed product.',
    )
    args = parser.parse_args(argv)

    files = _iter_test2a_files(args.products)
    if not files:
        print(f'No test_2a xlsx files found in {TEST_2A_DIR}')
        return

    print('test_2b2: CLI adjustment_mul comparison against main_mink')
    print(f'Input : {TEST_2A_DIR}')
    print(f'Main  : {MAIN_MINK_DIR}')
    print(f'Workers: {args.workers}')
    print(f'Tolerance: {args.tolerance}')

    results: list[ProductComparison] = []
    errors: list[tuple[str, Exception]] = []

    if args.workers <= 1 or len(files) <= 1:
        for path in files:
            try:
                result = compare_product(path, args.tolerance, args.examples)
                results.append(result)
                _print_result(result)
            except Exception as exc:
                errors.append((path.stem, exc))
                print(f'  {path.stem}: ERROR {exc}')
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            futures = {
                pool.submit(compare_product, path, args.tolerance, args.examples): path
                for path in files
            }
            for future in as_completed(futures):
                path = futures[future]
                try:
                    result = future.result()
                    results.append(result)
                    _print_result(result)
                except Exception as exc:
                    errors.append((path.stem, exc))
                    print(f'  {path.stem}: ERROR {exc}')

    failed = [result for result in results if not result.ok]
    print()
    print(f'Done! {len(results) - len(failed)} OK, {len(failed)} failed, {len(errors)} errors, {len(files)} total')

    if failed:
        print('\nFailed products:')
        for result in sorted(failed, key=lambda r: r.product)[:20]:
            print(
                f'  {result.product}: missing={result.missing_trade_time}, '
                f'diff={result.diff_rows}, max_abs_diff={result.max_abs_diff}'
            )

    if errors:
        print('\nErrored products:')
        for prod, exc in errors[:20]:
            print(f'  {prod}: {exc}')

    if failed or errors:
        sys.exit(1)


if __name__ == '__main__':
    main()
