import argparse
import os
import shutil
import sys
import tempfile
import zipfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, numbers
from openpyxl.utils.datetime import from_excel

from tests.calc import (
    FORMULA_FILL,
    HEADER_FILL,
    REMARK_FILL,
    REMARK_FONT,
    TEST_2A_DIR,
    TEST_3A_DIR,
    remark_height,
)

"""
test_3a_next_return.py
----------------------
Generate an Excel workbook that manually verifies RF-period open-to-open returns
from the full test_2a MAIN continuous-main sequence.

Input:
  data/test/test_2a/{prod}.xlsx, sheet MAIN, cached values after WPS/Excel save.

Output:
  data/test/test_3a/{prod}.xlsx

Workbook layout:
  MAIN_OPEN:
    Full MAIN rows, but only the columns needed for open-to-open return audit:
    trading_day, trade_time, open_price, adjustment_mul.

  OPEN_TO_OPEN_RF:
    B2 is RF_MINUTES.  A4 contains one dynamic-array Excel formula that spills
    all rows for the selected RF on 1min data frequency ($F=1min):
      open_adjusted = open_price * adjustment_mul
      return = next_open_adjusted / current_open_adjusted - 1

Python copies source values and writes formula text.  It does not hardcode
Python-calculated return results into Excel.
"""


DEFAULT_WORKERS = max(1, min(4, os.cpu_count() or 1))

NS_SHEET = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('', NS_SHEET)

MAIN_SHEET = 'MAIN'
MAIN_OPEN_SHEET = 'MAIN_OPEN'
RETURN_SHEET = 'OPEN_TO_OPEN_RF'

MAIN_HEADER_ROW = 2
MAIN_DATA_START_ROW = 3
MAIN_OPEN_HEADER_ROW = 2
MAIN_OPEN_DATA_START_ROW = 3
RETURN_CONFIG_ROW = 2
RETURN_HEADER_ROW = 3
RETURN_FORMULA_ROW = 4

SOURCE_HEADERS = [
    'trading_day',
    'trade_time',
    'open_price',
    'adjustment_mul',
]

RETURN_HEADERS = [
    'trading_day',
    'trade_time',
    'open_price',
    'adjustment_mul',
    'open_price_adjusted_formula',
    'next_trade_time',
    'next_open_price_adjusted_formula',
    'next_open_to_open_adjusted_rf',
]

DATE_FORMAT = 'yyyy-mm-dd'
DATETIME_FORMAT = 'yyyy-mm-dd hh:mm:ss'


def _normalise_header(value) -> str:
    return '' if value is None else str(value)


def _col_index(headers: list, name: str) -> int:
    normalised = [_normalise_header(h) for h in headers]
    try:
        return normalised.index(name) + 1
    except ValueError as exc:
        raise RuntimeError(f'MAIN missing required column: {name}') from exc


def _normalise_excel_date(value):
    if hasattr(value, 'to_pydatetime'):
        return value.to_pydatetime()
    if isinstance(value, (int, float)):
        return from_excel(value)
    return value


def _normalise_value(value):
    if hasattr(value, 'to_pydatetime'):
        return value.to_pydatetime()
    return value


def _iter_test2a_files(products: Iterable[str] | None = None) -> list[Path]:
    if products:
        return [TEST_2A_DIR / f'{prod}.xlsx' for prod in products]
    return sorted(TEST_2A_DIR.glob('*.xlsx'))


def _style_header_row(ws, row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row, col)
        cell.fill = HEADER_FILL
        cell.font = REMARK_FONT


def _write_remark(ws, text: str, max_col: int) -> None:
    ws.cell(1, 1, text)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.row_dimensions[1].height = remark_height(text)
    cell = ws.cell(1, 1)
    cell.fill = REMARK_FILL
    cell.font = REMARK_FONT
    cell.alignment = Alignment(wrap_text=True, vertical='top')


def _set_main_open_layout(ws) -> None:
    ws.freeze_panes = 'A3'
    widths = {'A': 14, 'B': 22, 'C': 14, 'D': 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _set_return_layout(ws) -> None:
    ws.freeze_panes = 'A4'
    widths = {
        'A': 14,
        'B': 22,
        'C': 14,
        'D': 16,
        'E': 26,
        'F': 22,
        'G': 30,
        'H': 28,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions['A'].number_format = DATE_FORMAT
    ws.column_dimensions['B'].number_format = DATETIME_FORMAT
    ws.column_dimensions['F'].number_format = DATETIME_FORMAT
    ws.column_dimensions['H'].number_format = numbers.FORMAT_PERCENTAGE_00


def _load_main_open_rows(src_path: Path) -> tuple[list[tuple], int]:
    value_wb = load_workbook(src_path, read_only=True, data_only=True)
    try:
        if MAIN_SHEET not in value_wb.sheetnames:
            raise RuntimeError(f'{src_path.name} has no MAIN sheet')

        value_ws = value_wb[MAIN_SHEET]
        headers = [
            _normalise_header(value_ws.cell(MAIN_HEADER_ROW, col).value)
            for col in range(1, value_ws.max_column + 1)
        ]
        if not headers or not headers[0] or value_ws.cell(MAIN_DATA_START_ROW, 1).value is None:
            raise RuntimeError(
                f'{src_path.name} MAIN has no cached values. '
                'Open test_2a workbook in WPS/Excel, let formulas calculate, save it, then rerun test_3a.'
            )

        source_indexes = [_col_index(headers, name) for name in SOURCE_HEADERS]
        rows: list[tuple] = []
        for row in value_ws.iter_rows(
            min_row=MAIN_DATA_START_ROW,
            max_row=value_ws.max_row,
            min_col=1,
            max_col=value_ws.max_column,
            values_only=True,
        ):
            trading_day = _normalise_excel_date(row[source_indexes[0] - 1])
            trade_time = _normalise_excel_date(row[source_indexes[1] - 1])
            open_price = _normalise_value(row[source_indexes[2] - 1])
            adjustment_mul = _normalise_value(row[source_indexes[3] - 1])
            rows.append((trading_day, trade_time, open_price, adjustment_mul))

        return rows, len(rows)
    finally:
        value_wb.close()


def _write_main_open_sheet(wb: Workbook, rows: list[tuple]) -> None:
    ws = wb.create_sheet(MAIN_OPEN_SHEET)
    _set_main_open_layout(ws)
    _write_remark(
        ws,
        'Full MAIN sequence from test_2a, reduced to the fields needed for open-to-open return audit.',
        len(SOURCE_HEADERS),
    )
    ws.append(SOURCE_HEADERS)
    _style_header_row(ws, MAIN_OPEN_HEADER_ROW, len(SOURCE_HEADERS))

    for row in rows:
        ws.append(row)

    for cell in ws.iter_rows(min_row=MAIN_OPEN_DATA_START_ROW, min_col=1, max_col=1):
        cell[0].number_format = DATE_FORMAT
    for cell in ws.iter_rows(min_row=MAIN_OPEN_DATA_START_ROW, min_col=2, max_col=2):
        cell[0].number_format = DATETIME_FORMAT


def _build_return_formula(main_rows: int) -> str:
    source_start = MAIN_OPEN_DATA_START_ROW
    source_end = MAIN_OPEN_DATA_START_ROW + main_rows - 1
    return (
        '=LET('
        'rf,MAX(1,INT($B$2)),'
        f'n,ROWS(MAIN_OPEN!A{source_start}:A{source_end}),'
        'seq,SEQUENCE(n),'
        'cur_row,seq+2,'
        'next_row,seq+rf+2,'
        'has_next,seq+rf<=n,'
        'cur_open,INDEX(MAIN_OPEN!C:C,cur_row)*INDEX(MAIN_OPEN!D:D,cur_row),'
        'next_open,IF(has_next,INDEX(MAIN_OPEN!C:C,next_row)*INDEX(MAIN_OPEN!D:D,next_row),""),'
        'HSTACK('
        'INDEX(MAIN_OPEN!A:A,cur_row),'
        'INDEX(MAIN_OPEN!B:B,cur_row),'
        'INDEX(MAIN_OPEN!C:C,cur_row),'
        'INDEX(MAIN_OPEN!D:D,cur_row),'
        'cur_open,'
        'IF(has_next,INDEX(MAIN_OPEN!B:B,next_row),""),'
        'next_open,'
        'IF(has_next,next_open/cur_open-1,"")'
        '))'
    )


def _write_return_sheet(wb: Workbook, main_rows: int) -> None:
    ws = wb.create_sheet(RETURN_SHEET)
    _set_return_layout(ws)
    _write_remark(
        ws,
        'Set RF_MINUTES in B2. A4 spills the full RF-period open-to-open adjusted return table on $F=1min data.',
        len(RETURN_HEADERS),
    )

    ws.cell(RETURN_CONFIG_ROW, 1, 'RF_MINUTES')
    ws.cell(RETURN_CONFIG_ROW, 2, 1)
    ws.cell(RETURN_CONFIG_ROW, 4, 'DATA_FREQ')
    ws.cell(RETURN_CONFIG_ROW, 5, '1min')
    for col in (1, 4):
        ws.cell(RETURN_CONFIG_ROW, col).fill = HEADER_FILL
        ws.cell(RETURN_CONFIG_ROW, col).font = REMARK_FONT
    ws.cell(RETURN_CONFIG_ROW, 2).fill = FORMULA_FILL

    ws.append(RETURN_HEADERS)
    _style_header_row(ws, RETURN_HEADER_ROW, len(RETURN_HEADERS))

    formula_cell = ws.cell(RETURN_FORMULA_ROW, 1, _build_return_formula(main_rows))
    formula_cell.fill = FORMULA_FILL
    formula_cell.number_format = DATE_FORMAT


def _patch_dynamic_array_formula(xlsx_path: Path, main_rows: int) -> None:
    max_output_row = RETURN_FORMULA_ROW + main_rows - 1
    array_ref = f'A{RETURN_FORMULA_ROW}:H{max_output_row}'
    tmp_fd, tmp_path = tempfile.mkstemp(suffix='.xlsx')
    os.close(tmp_fd)
    tmp = Path(tmp_path)
    try:
        with zipfile.ZipFile(xlsx_path, 'r') as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
            target_sheet = 'xl/worksheets/sheet2.xml'
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == target_sheet:
                    root = ET.fromstring(data)
                    cell = root.find(f".//{{{NS_SHEET}}}c[@r='A{RETURN_FORMULA_ROW}']")
                    if cell is not None:
                        formula = cell.find(f'{{{NS_SHEET}}}f')
                        if formula is not None:
                            formula.set('t', 'array')
                            formula.set('ref', array_ref)
                            formula.set('ca', '1')
                            formula.set('aca', '1')
                    data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                zout.writestr(item, data)
        shutil.move(str(tmp), xlsx_path)
    finally:
        if tmp.exists():
            tmp.unlink()


def process_product(test2a_xlsx: Path) -> int:
    prod = test2a_xlsx.stem
    if not test2a_xlsx.exists():
        print(f'  test_2a file not found: {test2a_xlsx.name}, skipping')
        return 0

    TEST_3A_DIR.mkdir(parents=True, exist_ok=True)
    out_path = TEST_3A_DIR / f'{prod}.xlsx'

    print(f'  Processing: {prod}...', end='', flush=True)
    rows, main_rows = _load_main_open_rows(test2a_xlsx)

    wb = Workbook()
    active = wb.active
    if active is not None:
        wb.remove(active)

    _write_main_open_sheet(wb, rows)
    _write_return_sheet(wb, main_rows)
    wb.save(out_path)
    wb.close()
    _patch_dynamic_array_formula(out_path, main_rows)

    print(f' OK ({main_rows} MAIN rows)')
    return main_rows


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description='Generate test_3a RF open-to-open return workbooks.')
    parser.add_argument('products', nargs='*', help='Optional product ids, e.g. A RB IF')
    parser.add_argument(
        '--workers',
        type=int,
        default=int(os.environ.get('TEST_3A_WORKERS', DEFAULT_WORKERS)),
        help='Number of product workbooks to generate in parallel.',
    )
    args = parser.parse_args(argv)

    files = _iter_test2a_files(args.products)
    if not files:
        print(f'No test_2a xlsx files found in {TEST_2A_DIR}')
        return

    print('test_3a: Excel RF NEXT_OPEN_TO_OPEN_ADJUSTED audit source')
    print(f'Input : {TEST_2A_DIR}')
    print(f'Output: {TEST_3A_DIR}')
    print(f'Workers: {args.workers}')

    errors: list[tuple[str, Exception]] = []
    if args.workers <= 1 or len(files) <= 1:
        for path in files:
            try:
                process_product(path)
            except Exception as exc:
                errors.append((path.stem, exc))
                print(f' ERROR {exc}')
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            futures = {pool.submit(process_product, path): path for path in files}
            for future in as_completed(futures):
                path = futures[future]
                try:
                    future.result()
                except Exception as exc:
                    errors.append((path.stem, exc))
                    print(f'  {path.stem}: ERROR {exc}')

    if errors:
        print(f'\nFailed products: {len(errors)}')
        for prod, exc in errors[:20]:
            print(f'  {prod}: {exc}')
        sys.exit(1)


def test_return_formula_workbook_for_a_product():
    main_rows = process_product(TEST_2A_DIR / 'A.xlsx')

    out_path = TEST_3A_DIR / 'A.xlsx'
    wb = load_workbook(out_path, read_only=False, data_only=False)
    try:
        assert wb.sheetnames == [MAIN_OPEN_SHEET, RETURN_SHEET]
        main_ws = wb[MAIN_OPEN_SHEET]
        return_ws = wb[RETURN_SHEET]

        headers = [main_ws.cell(MAIN_OPEN_HEADER_ROW, col).value for col in range(1, 5)]
        assert headers == SOURCE_HEADERS
        assert main_rows > 1000
        assert main_ws.max_column == 4
        assert str(next(iter(main_ws.merged_cells.ranges))) == 'A1:D1'
        assert str(next(iter(return_ws.merged_cells.ranges))) == 'A1:H1'
        assert main_ws.cell(MAIN_OPEN_DATA_START_ROW, 1).number_format == DATE_FORMAT
        assert main_ws.cell(MAIN_OPEN_DATA_START_ROW, 2).number_format == DATETIME_FORMAT

        assert return_ws.cell(RETURN_CONFIG_ROW, 1).value == 'RF_MINUTES'
        assert return_ws.cell(RETURN_CONFIG_ROW, 2).value == 1
        assert return_ws.cell(RETURN_HEADER_ROW, 8).value == 'next_open_to_open_adjusted_rf'

        array_formula = return_ws.cell(RETURN_FORMULA_ROW, 1).value
        formula = getattr(array_formula, 'text', array_formula)
        assert formula.startswith('=LET(')
        assert '$B$2' in formula
        assert 'SEQUENCE(n)' in formula
        assert 'has_next,seq+rf<=n' in formula
        assert 'IF(has_next,next_open/cur_open-1,"")' in formula
        assert getattr(array_formula, 'ref', None) == f'A4:H{RETURN_FORMULA_ROW + main_rows - 1}'
        assert return_ws.max_row == RETURN_FORMULA_ROW
    finally:
        wb.close()


if __name__ == '__main__':
    main()
