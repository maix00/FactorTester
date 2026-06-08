"""
test_2b_compare_main_mink.py
----------------------------
Generate a lightweight Excel check for test_2a MAIN adjustment factors.

This script does not copy all MAIN data.  It only extracts:
  - trade_time
  - adjustment_mul

Rows are matched by trade_time.  The output workbook uses Excel formulas to
verify that test_2a's adjustment_mul equals main_mink's adjustment_mul.

Prerequisite:
  test_2a/{prod}.xlsx must have been opened, calculated, and saved by WPS/Excel.
  openpyxl can read cached formula values, but it cannot calculate test_2a's
  dynamic-array formulas.
"""

from __future__ import annotations

import argparse
import os
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path
from typing import Iterable

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, numbers
from openpyxl.utils import get_column_letter, range_boundaries

from tests.calc import (
    HEADER_FILL,
    MAIN_MINK_DIR,
    REMARK_FILL,
    REMARK_FONT,
    TEST_2A_DIR,
    TEST_2B_DIR,
    remark_height,
)


DEFAULT_WORKERS = max(1, min(4, os.cpu_count() or 1))
TOLERANCE = 1e-8

MAIN_SHEET = 'MAIN'
TEST2A_ADJ_SHEET = 'TEST_2A_ADJ'
MAIN_MINK_ADJ_SHEET = 'MAIN_MINK_ADJ'
COMPARE_SHEET = 'COMPARE'

MAIN_HEADER_ROW = 2
MAIN_DATA_START_ROW = 3
MINK_DATA_START_ROW = 2

TRADE_TIME_COL = 'trade_time'
ADJ_MUL_COL = 'adjustment_mul'
KEY_HEADER = '_trade_time_key'


def _excel_ref_bounds(ws, cell_ref: str) -> tuple[int | None, int | None, int | None, int | None]:
    value = ws[cell_ref].value
    ref = getattr(value, 'ref', None)
    if ref:
        return range_boundaries(ref)  # type: ignore[return-value]
    return (ws[cell_ref].column, ws[cell_ref].row, ws.max_column, ws.max_row)


def _find_main_mink_file(prod: str) -> Path | None:
    candidates = [
        p for p in MAIN_MINK_DIR.glob(f'{prod}.*.parquet')
        if p.name.split('.', 1)[0] == prod
    ]
    if not candidates:
        return None
    if len(candidates) > 1:
        names = ', '.join(p.name for p in sorted(candidates))
        raise RuntimeError(f'multiple main_mink files found for {prod}: {names}')
    return candidates[0]


def _normalise_value(value):
    if pd.isna(value):
        return None
    if hasattr(value, 'to_pydatetime'):
        return value.to_pydatetime()
    return value


def _styled_cell(ws, value, *, fill=None, font=None, alignment=None):
    cell = WriteOnlyCell(ws, value=value)
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    return cell


def _write_header(ws, headers: list[str]) -> None:
    ws.append([
        _styled_cell(ws, h, fill=HEADER_FILL, font=REMARK_FONT)
        for h in headers
    ])


def _col_index(headers: list, name: str) -> int:
    try:
        return [str(h) if h is not None else '' for h in headers].index(name) + 1
    except ValueError as exc:
        raise RuntimeError(f'MAIN missing required column: {name}') from exc


def _write_test2a_adj(src_path: Path, ws_out) -> int:
    formula_wb = load_workbook(src_path, read_only=True, data_only=False)
    value_wb = load_workbook(src_path, read_only=True, data_only=True)
    try:
        if MAIN_SHEET not in formula_wb.sheetnames or MAIN_SHEET not in value_wb.sheetnames:
            raise RuntimeError(f'{src_path.name} has no MAIN sheet')

        formula_ws = formula_wb[MAIN_SHEET]
        value_ws = value_wb[MAIN_SHEET]
        _min_col, _min_row, max_col, max_row = _excel_ref_bounds(formula_ws, 'A3')

        headers = [value_ws.cell(MAIN_HEADER_ROW, c).value for c in range(1, max_col + 1)]
        if not headers or headers[0] is None or value_ws.cell(MAIN_DATA_START_ROW, 1).value is None:
            raise RuntimeError(
                f'{src_path.name} MAIN has no cached calculated values. '
                'Open test_2a workbook in WPS/Excel, let formulas calculate, save it, then rerun test_2b.'
            )

        trade_time_col = _col_index(headers, TRADE_TIME_COL)
        adj_col = _col_index(headers, ADJ_MUL_COL)

        remark = 'Cached values extracted from test_2a MAIN: trade_time + adjustment_mul only.'
        ws_out.append([
            _styled_cell(
                ws_out, remark, fill=REMARK_FILL, font=REMARK_FONT,
                alignment=Alignment(wrap_text=True, vertical='top')
            )
        ])
        ws_out.row_dimensions[1].height = remark_height(remark)
        _write_header(ws_out, [TRADE_TIME_COL, ADJ_MUL_COL, KEY_HEADER])

        out_row = 3
        for row in value_ws.iter_rows(
            min_row=MAIN_DATA_START_ROW,
            max_row=max_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        ):
            trade_time = row[trade_time_col - 1]
            adj_mul = row[adj_col - 1]
            key_formula = f'=TEXT(A{out_row},"yyyy-mm-dd hh:mm:ss")'
            time_cell = _styled_cell(
                ws_out, trade_time,
                alignment=Alignment(horizontal='right'),
            )
            time_cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            ws_out.append([time_cell, adj_mul, key_formula])
            out_row += 1

        return out_row - 3
    finally:
        formula_wb.close()
        value_wb.close()


def _write_main_mink_adj(mink_path: Path, ws_out) -> int:
    df = pd.read_parquet(mink_path, columns=[TRADE_TIME_COL, ADJ_MUL_COL])

    _write_header(ws_out, [TRADE_TIME_COL, ADJ_MUL_COL, KEY_HEADER])
    for out_row, (trade_time, adj_mul) in enumerate(
        df.itertuples(index=False, name=None), start=MINK_DATA_START_ROW
    ):
        key_formula = f'=TEXT(A{out_row},"yyyy-mm-dd hh:mm:ss")'
        time_cell = _styled_cell(
            ws_out, _normalise_value(trade_time),
            alignment=Alignment(horizontal='right'),
        )
        time_cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        ws_out.append([time_cell, _normalise_value(adj_mul), key_formula])

    return len(df)


def _write_compare(ws_out, test2a_rows: int, mink_rows: int) -> None:
    compare_start_row = 12
    compare_end_row = compare_start_row + test2a_rows - 1

    remark = (
        'Excel formula comparison: match trade_time and compare adjustment_mul '
        'between TEST_2A_ADJ and MAIN_MINK_ADJ.'
    )
    ws_out.append([
        _styled_cell(
            ws_out, remark, fill=REMARK_FILL, font=REMARK_FONT,
            alignment=Alignment(wrap_text=True, vertical='top')
        )
    ])
    ws_out.row_dimensions[1].height = remark_height(remark)
    ws_out.append(['tolerance', TOLERANCE])
    ws_out.append([])
    ws_out.append(['test_2a_rows', test2a_rows])
    ws_out.append(['main_mink_rows', mink_rows])
    ws_out.append(['missing_trade_time', f'=COUNTIF(F{compare_start_row}:F{compare_end_row},"MISSING")'])
    ws_out.append(['diff_rows', f'=COUNTIF(F{compare_start_row}:F{compare_end_row},"FAIL")'])
    ws_out.append(['max_abs_diff', f'=MAX(E{compare_start_row}:E{compare_end_row})'])
    ws_out.append(['overall', '=IF(AND(B6=0,B7=0),"PASS","FAIL")'])
    ws_out.append([])

    _write_header(
        ws_out,
        [
            '_test2a_row',
            TRADE_TIME_COL,
            'test_2a_adjustment_mul',
            'main_mink_adjustment_mul',
            'abs_diff',
            'status',
        ],
    )

    mink_key_col = 'C'
    mink_adj_col = 'B'
    for row in range(compare_start_row, compare_end_row + 1):
        source_row = f'=ROW()-{compare_start_row - 3}'
        trade_time = _styled_cell(
            ws_out,
            f'=INDEX(TEST_2A_ADJ!A:A,A{row})',
            alignment=Alignment(horizontal='right'),
        )
        trade_time.number_format = 'yyyy-mm-dd hh:mm:ss'
        test2a_adj = f'=INDEX(TEST_2A_ADJ!B:B,A{row})'
        key = f'TEXT(B{row},"yyyy-mm-dd hh:mm:ss")'
        mink_row = f'MATCH({key},MAIN_MINK_ADJ!${mink_key_col}:${mink_key_col},0)'
        mink_adj = f'=IFERROR(INDEX(MAIN_MINK_ADJ!${mink_adj_col}:${mink_adj_col},{mink_row}),"")'
        abs_diff = f'=IF(D{row}="","",ABS(C{row}-D{row}))'
        status = f'=IF(D{row}="","MISSING",IF(E{row}<=$B$2,"PASS","FAIL"))'
        ws_out.append([source_row, trade_time, test2a_adj, mink_adj, abs_diff, status])


def process_product(test2a_xlsx: Path) -> None:
    prod = test2a_xlsx.stem
    if not test2a_xlsx.exists():
        print(f'  test_2a file not found: {test2a_xlsx.name}, skipping')
        return

    mink_path = _find_main_mink_file(prod)
    if mink_path is None:
        print(f'  main_mink file not found for {prod}, skipping')
        return

    TEST_2B_DIR.mkdir(parents=True, exist_ok=True)
    out_path = TEST_2B_DIR / f'{prod}.xlsx'

    print(f'  Processing: {prod}...', end='', flush=True)
    wb = Workbook(write_only=True)
    test2a_ws = wb.create_sheet(TEST2A_ADJ_SHEET)
    mink_ws = wb.create_sheet(MAIN_MINK_ADJ_SHEET)
    compare_ws = wb.create_sheet(COMPARE_SHEET)

    test2a_rows = _write_test2a_adj(test2a_xlsx, test2a_ws)
    mink_rows = _write_main_mink_adj(mink_path, mink_ws)
    _write_compare(compare_ws, test2a_rows=test2a_rows, mink_rows=mink_rows)

    wb.save(out_path)
    print(f' OK ({test2a_rows} test_2a rows, {mink_rows} main_mink rows)')


def _iter_test2a_files(products: Iterable[str] | None = None) -> list[Path]:
    if products:
        return [TEST_2A_DIR / f'{prod}.xlsx' for prod in products]
    return sorted(TEST_2A_DIR.glob('*.xlsx'))


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description='Generate test_2b adjustment factor comparisons.')
    parser.add_argument('products', nargs='*', help='Optional product ids, e.g. A RB IF')
    parser.add_argument(
        '--workers',
        type=int,
        default=int(os.environ.get('TEST_2B_WORKERS', DEFAULT_WORKERS)),
        help='Number of product workbooks to generate in parallel.',
    )
    args = parser.parse_args(argv)

    files = _iter_test2a_files(args.products)
    if not files:
        print(f'No test_2a xlsx files found in {TEST_2A_DIR}')
        return

    print('test_2b: Excel adjustment_mul comparison against main_mink')
    print(f'Input : {TEST_2A_DIR}')
    print(f'Output: {TEST_2B_DIR}')
    print(f'Workers: {args.workers}')

    errors: list[tuple[str, Exception]] = []
    if args.workers <= 1 or len(files) <= 1:
        for path in files:
            try:
                process_product(path)
            except Exception as exc:
                errors.append((path.stem, exc))
                print(f' ERROR {exc}')
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            futures = {pool.submit(process_product, path): path for path in files}
            for future in as_completed(futures):
                path = futures[future]
                try:
                    future.result()
                except Exception as exc:
                    errors.append((path.stem, exc))
                    print(f'  {path.stem}: ERROR {exc}')

    if errors:
        print(f'\nFailed products: {len(errors)}')
        for prod, exc in errors[:20]:
            print(f'  {prod}: {exc}')
        sys.exit(1)


if __name__ == '__main__':
    main()
