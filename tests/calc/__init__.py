"""
tests/calc 包初始化。

职责边界：
1. 集中声明 calc 测试允许读取的数据源路径。
2. 集中声明 calc 测试输出目录。
3. 提供跨测试共用的轻量配置与 Excel 样式常量。

本文件不记录单个测试脚本的实现细节；例如 test_2a 的动态数组、
XML patch、WPS 显示兼容等说明，应留在 test_2a 脚本自身。
"""

from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill

from scripts.data_dir import DATA_DIR as _ROOT_DATA_DIR

# ============================================================
# 数据源路径（只读）
# ============================================================

# 核心计算输入：test_0 / test_1 / test_2a 只能从这些源推导结果。
WIND_MAPPING_PATH = Path(_ROOT_DATA_DIR) / 'wind_mapping.parquet'
MIN_DATA_DIR = Path(_ROOT_DATA_DIR) / 'data_mink_product'

# 对账输入：只允许 test_2b 用于校验 test_2a 的 MAIN 结果，不允许 test_2a 读取。
MAIN_MINK_DIR = Path(_ROOT_DATA_DIR) / 'main_mink'

# ============================================================
# Test 输出路径
# ============================================================

# test_0: 统计分析结果
TEST_0_DIR = Path(_ROOT_DATA_DIR) / 'test' / 'test_0'
TEST_0_PRODUCTS_XLSX = TEST_0_DIR / '_products.xlsx'

# test_1: 截断数据导出（每个品种一个 Excel，直接放在此目录下）
TEST_1_DIR = Path(_ROOT_DATA_DIR) / 'test' / 'test_1'

TEST_2A_DIR = Path(_ROOT_DATA_DIR) / 'test' / 'test_2a'

# test_2b / test_2c (预留给后续)
TEST_2B_DIR = Path(_ROOT_DATA_DIR) / 'test' / 'test_2b'
TEST_2C_DIR = Path(_ROOT_DATA_DIR) / 'test' / 'test_2c'

# test_3a: 从 test_2a MAIN 全量主力序列验证下一期收益率公式
TEST_3A_DIR = Path(_ROOT_DATA_DIR) / 'test' / 'test_3a'
TEST_3B_DIR = Path(_ROOT_DATA_DIR) / 'test' / 'test_3b'

# ============================================================
# 全局配置
# ============================================================

WINDOW_ROWS = 1  # test_1 合约数据主力区间前后额外保留的分钟行数

# 数据源文档说明（供 agent 和 human 查阅）
_DATA_SOURCES_DOC = f"""
=== calc 包数据源白名单 ===
核心计算输入（test_0 / test_1 / test_2a）:
1. WIND_MAPPING_PATH = \"{WIND_MAPPING_PATH}\"
   → Wind 主力合约映射表，字段: S_INFO_WINDCODE, FS_MAPPING_WINDCODE, STARTDATE, ENDDATE
2. MIN_DATA_DIR = \"{MIN_DATA_DIR}\"
   → 分钟行情数据，文件命名: {{交易所}}|F|{{品种}}|{{合约}}.parquet

对账输入（仅 test_2b）:
3. MAIN_MINK_DIR = \"{MAIN_MINK_DIR}\"
   → 已生成的分钟主力连续序列，只允许用于和 test_2a 的 MAIN 做结果对账。

禁止使用其他数据源（main_dayk, 自定义CSV/Parquet/SQLite等）。
"""

# wind_mapping 交易所 → 分钟文件 交易所
WIND_EXCH_TO_MIN_EXCH: dict[str, str] = {
    'SHF': 'SHFE',
    'DCE': 'DCE',
    'CZC': 'CZCE',
    'INE': 'INE',
    'CFE': 'CFFEX',
    'GFE': 'GFEX',
}

# ============================================================
# 共享 Excel 样式常量
# ============================================================

FORMULA_FILL = PatternFill(start_color='FFFFF2CC', end_color='FFFFF2CC', fill_type='solid')
HEADER_FILL = PatternFill(start_color='FFD9E1F2', end_color='FFD9E1F2', fill_type='solid')

# Row 1 remark 样式（test_1 / test_2a 共用）
REMARK_FILL = PatternFill(start_color='FFFCE4D6', end_color='FFFCE4D6', fill_type='solid')
REMARK_FONT = Font(bold=True)        # 也用于表头
REMARK_ALIGNMENT = Alignment(wrap_text=True, vertical='top')

# test_1 comment 样式
COMMENT_FONT = Font(italic=True, size=10)
COMMENT_ALIGNMENT = Alignment(wrap_text=True)  # 与 test_1 一致（无 vertical='top'）

# test_1 → 5 行 remark → 80pt；以此基准
REMARK_ROW_HEIGHT_PER_LINE = 16


def remark_height(text: str) -> float:
    """按行数动态计算 row 1 remark 高度，最小 45pt，最大 200pt"""
    lines = text.count('\n') + 1
    return max(45, min(200, lines * REMARK_ROW_HEIGHT_PER_LINE))
