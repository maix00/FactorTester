# pyright: reportMissingImports=false, reportMissingModuleSource=false
"""
test_4b_rolling_mean.py
-----------------------
验证 RollingOp('rolling_mean') 的行为。

从 test_2a MAIN sheet 读取主力连续序列（含复权因子 adj_mul），
用后端 rolling_mean 分别计算 10min / 5bar / 1day / 1d30min 四种窗口，
写入 Excel 让用户手工验算。
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment
from openpyxl.utils.datetime import from_excel

from sources.LocalCNFutures.CNFutures import CNFutures
from tools.data.types.DataIndex import finest_index
from tools.data.types.DataColumn import DataColumn
from tools.data.types.DataFreq import DataFreq
from tools.factors.FactorExpr import ColumnRef, EvaluateContext
from tools.factors.FactorFamily import FactorFamily
from tools.factors.FactorTester import FactorTester
from tools.factors.expr.leaf import ConstExpr
from tools.factors.expr.rolling import RollingOp, _resolve_windows

from tests.calc import (
    HEADER_FILL,
    MAIN_MINK_DIR,
    COMMENT_FONT,
    REMARK_FILL,
    REMARK_FONT,
    TEST_2A_DIR,
    TEST_4B_DIR,
    remark_height,
)

PRICE_SHEET = 'PRICE'
BACKEND_SHEET = 'BACKEND'
COMPARE_SHEET = 'COMPARE'

DT_FMT = 'yyyy-mm-dd hh:mm:ss'
DATE_FMT = 'yyyy-mm-dd'
NUM_FMT = '0.0000000000'
ROW_FMT = '0'

WINDOWS = [
    ('ROLLING_MEAN_10MIN', '10min'),
    ('ROLLING_MEAN_5BAR', 5),
    ('ROLLING_MEAN_1DAY', '1day'),
    ('ROLLING_MEAN_1DAY_30MIN', '1d30min'),
]

COL_TRADING_DAY = 0
COL_TRADE_TIME = 1
COL_OPEN_PRICE = 6
COL_ADJ_MUL = 21

DATA_ROW_OFFSET = 2
BACKEND_DATA_START_ROW = 3
COMPARE_START_ROW = 14

NS_SHEET = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ET.register_namespace('', NS_SHEET)
ET.register_namespace('r', NS_R)


class _OpenAdjustedFactor(FactorFamily):

    @staticmethod
    def factor_expr():
        return ColumnRef(DataColumn.OPEN_ADJUSTED)


def _styled_cell(ws, value, *, fill=None, font=None, alignment=None, number_format=None):
    cell = WriteOnlyCell(ws, value=value)
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if number_format is not None:
        cell.number_format = number_format
    return cell


def _write_remark(ws, text: str) -> None:
    ws.append([
        _styled_cell(ws, text, fill=REMARK_FILL, font=COMMENT_FONT,
                     alignment=Alignment(wrap_text=True, vertical='top'))
    ])
    ws.row_dimensions[1].height = remark_height(text)


def _write_header(ws, headers: list[str]) -> None:
    ws.append([
        _styled_cell(ws, h, fill=HEADER_FILL, font=REMARK_FONT) for h in headers
    ])


def _product_name(product_id: str) -> str:
    if '.' in product_id:
        return product_id
    candidates = [
        path.stem
        for path in MAIN_MINK_DIR.glob(f'{product_id}.*.parquet')
        if path.name.split('.', 1)[0] == product_id
    ]
    if not candidates:
        raise RuntimeError(f'Cannot infer exchange for product {product_id}; pass full name like A.DCE')
    if len(candidates) > 1:
        raise RuntimeError(f'Multiple products match {product_id}: {", ".join(sorted(candidates))}')
    return candidates[0]


def _normalise_excel_datetime(value):
    if hasattr(value, 'to_pydatetime'):
        return value.to_pydatetime()
    if isinstance(value, (int, float)):
        return from_excel(value)
    return value


def _normalise_series_index(series: pd.Series) -> pd.Series:
    index = finest_index(series.index) if isinstance(series.index, pd.MultiIndex) else series.index
    ts_index = pd.DatetimeIndex(index)
    if ts_index.tz is not None:
        ts_index = ts_index.tz_convert('Asia/Shanghai').tz_localize(None)
    series = series.copy()
    series.index = ts_index
    return series


def _window_bars(product: CNFutures) -> dict[str, int]:
    result: dict[str, int] = {}
    for name, window_val in WINDOWS:
        common, periods, _product_periods = _resolve_windows(
            window_val,
            DataFreq('1min'),
            [product],
        )
        if not common:
            raise RuntimeError(f'{product.name}: window {name} does not resolve to common bars')
        result[name] = int(periods)
    return result


def _read_test2a_main(prod_path: Path) -> tuple[list, list, list, list] | None:
    if not prod_path.exists():
        return None
    wb = load_workbook(prod_path, read_only=True, data_only=True)
    try:
        ws = wb['MAIN']
        trading_days = []
        times = []
        opens = []
        adjs = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            d = row[COL_TRADING_DAY] if len(row) > COL_TRADING_DAY else None
            t = row[COL_TRADE_TIME] if len(row) > COL_TRADE_TIME else None
            o = row[COL_OPEN_PRICE] if len(row) > COL_OPEN_PRICE else None
            a = row[COL_ADJ_MUL] if len(row) > COL_ADJ_MUL else None
            if t is None:
                break
            trading_days.append(_normalise_excel_datetime(d))
            times.append(_normalise_excel_datetime(t))
            opens.append(o)
            adjs.append(a)
        if not times:
            return None
        return trading_days, times, opens, adjs
    finally:
        wb.close()


def _open_adjusted_formula(last_row: int) -> str:
    return f'=C3:C{last_row}*D3:D{last_row}'


def _rolling_mean_formula(n_bars: int) -> str:
    min_periods = max(1, n_bars // 2)
    return (
        f'=LET(x,E3#,w,{n_bars},minp,{min_periods},'
        'MAP(SEQUENCE(ROWS(x)),LAMBDA(i,'
        'LET(s,MAX(1,i-w+1),c,i-s+1,'
        'IF(c<minp,"",AVERAGE(TAKE(DROP(x,s-1),c)))))))'
    )


def _compare_spill_formulas(nrows: int) -> list[str]:
    price_last_row = DATA_ROW_OFFSET + nrows
    backend_last_row = BACKEND_DATA_START_ROW + nrows - 1
    compare_last_row = COMPARE_START_ROW + nrows - 1
    return [
        f'=SEQUENCE({nrows})',
        f'=BACKEND!A{BACKEND_DATA_START_ROW}:A{backend_last_row}',
        f'=LET(r,BACKEND!B{BACKEND_DATA_START_ROW}:B{backend_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(r,PRICE!F3:F{price_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(b,C{COMPARE_START_ROW}:C{compare_last_row},x,D{COMPARE_START_ROW}:D{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb+ex,"",ABS(b-x)))',
        f'=LET(b,C{COMPARE_START_ROW}:C{compare_last_row},x,D{COMPARE_START_ROW}:D{compare_last_row},d,E{COMPARE_START_ROW}:E{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb*ex,"EMPTY",IF(eb+ex,"MISSING",IF(d<=$B$2,"PASS","FAIL"))))',
        f'=LET(r,BACKEND!C{BACKEND_DATA_START_ROW}:C{backend_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(r,PRICE!G3:G{price_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(b,G{COMPARE_START_ROW}:G{compare_last_row},x,H{COMPARE_START_ROW}:H{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb+ex,"",ABS(b-x)))',
        f'=LET(b,G{COMPARE_START_ROW}:G{compare_last_row},x,H{COMPARE_START_ROW}:H{compare_last_row},d,I{COMPARE_START_ROW}:I{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb*ex,"EMPTY",IF(eb+ex,"MISSING",IF(d<=$B$2,"PASS","FAIL"))))',
        f'=LET(r,BACKEND!D{BACKEND_DATA_START_ROW}:D{backend_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(r,PRICE!H3:H{price_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(b,K{COMPARE_START_ROW}:K{compare_last_row},x,L{COMPARE_START_ROW}:L{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb+ex,"",ABS(b-x)))',
        f'=LET(b,K{COMPARE_START_ROW}:K{compare_last_row},x,L{COMPARE_START_ROW}:L{compare_last_row},d,M{COMPARE_START_ROW}:M{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb*ex,"EMPTY",IF(eb+ex,"MISSING",IF(d<=$B$2,"PASS","FAIL"))))',
        f'=LET(r,BACKEND!E{BACKEND_DATA_START_ROW}:E{backend_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(r,PRICE!I3:I{price_last_row},IF(LEN(r&"")=0,"",r))',
        f'=LET(b,O{COMPARE_START_ROW}:O{compare_last_row},x,P{COMPARE_START_ROW}:P{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb+ex,"",ABS(b-x)))',
        f'=LET(b,O{COMPARE_START_ROW}:O{compare_last_row},x,P{COMPARE_START_ROW}:P{compare_last_row},d,Q{COMPARE_START_ROW}:Q{compare_last_row},eb,LEN(b&"")=0,ex,LEN(x&"")=0,IF(eb*ex,"EMPTY",IF(eb+ex,"MISSING",IF(d<=$B$2,"PASS","FAIL"))))',
    ]


def _write_price_sheet(ws, trading_days: list, times: list, opens: list, adjs: list, nrows: int, window_bars: dict[str, int]) -> None:
    _write_remark(ws, 'Price data from test_2a MAIN. '
                       'open_adj = open_price * adj_mul; rolling mean uses '
                       'min_periods = max(1, window_bars // 2), matching backend.')

    headers = [
        'trading_day',
        'trade_time',
        'open_price',
        'adj_mul',
        'open_adj',
        'ROLLING_MEAN_10MIN',
        'ROLLING_MEAN_5BAR',
        'ROLLING_MEAN_1DAY',
        'ROLLING_MEAN_1DAY_30MIN',
    ]
    _write_header(ws, headers)

    last_row = DATA_ROW_OFFSET + nrows
    for i in range(nrows):
        row_cells: list[Any] = [
            _styled_cell(ws, trading_days[i], number_format=DATE_FMT),
            _styled_cell(ws, times[i], number_format=DT_FMT),
            _styled_cell(ws, opens[i], number_format=NUM_FMT),
            _styled_cell(ws, adjs[i], number_format=NUM_FMT),
            _styled_cell(ws, _open_adjusted_formula(last_row), number_format=NUM_FMT) if i == 0 else None,
        ]
        for name, _window_val in WINDOWS:
            row_cells.append(
                _styled_cell(ws, _rolling_mean_formula(window_bars[name]), number_format=NUM_FMT)
                if i == 0 else None
            )
        ws.append(row_cells)


def _backend_rolling_mean(product: CNFutures, window_val) -> pd.Series:
    tester = FactorTester(products=[product], logger_file=False)
    factor = _OpenAdjustedFactor().get_factor(**{'$F': '1min', '$Rev': '0'})
    factor.evaluate([product])

    col_ref = ColumnRef(DataColumn.OPEN_ADJUSTED)
    rolling_expr = RollingOp('rolling_mean', ConstExpr(window_val), col_ref)

    freq = DataFreq('1min')
    ctx = EvaluateContext(products=[product], freq=freq)
    df = rolling_expr.evaluate(ctx=ctx)
    return _normalise_series_index(pd.Series(pd.to_numeric(df[product], errors='coerce')))


def _write_backend_sheet(ws, product: CNFutures, times: list) -> None:
    _write_remark(ws, 'Backend RollingOp(rolling_mean, window, ColumnRef(OPEN_ADJUSTED)) results.')

    headers = ['trade_time'] + [name for name, _ in WINDOWS]
    _write_header(ws, headers)

    rolling_series: dict[str, pd.Series] = {}
    for name, window_val in WINDOWS:
        try:
            rolling_series[name] = _backend_rolling_mean(product, window_val)
        except Exception as e:
            print(f'  WARNING: backend rolling_mean {name} failed: {e}')
            rolling_series[name] = pd.Series(dtype=float)

    for t in times:
        ts = pd.Timestamp(t)
        row_cells: list[Any] = [_styled_cell(ws, t, number_format=DT_FMT)]
        for name, _ in WINDOWS:
            s = rolling_series[name]
            if ts in s.index:
                val = s.loc[ts]
                row_cells.append(
                    _styled_cell(ws, float(val), number_format=NUM_FMT) if pd.notna(val) else ''
                )
            else:
                row_cells.append('')
        ws.append(row_cells)


def _write_compare_sheet(ws, nrows: int) -> None:
    _write_remark(ws, 'Compare PRICE Excel rolling_mean formulas vs BACKEND python rolling_mean.')

    compare_end_row = COMPARE_START_ROW + nrows - 1
    ws.append([
        _styled_cell(ws, 'TOLERANCE', fill=HEADER_FILL, font=REMARK_FONT),
        0.0001,
    ])
    ws.append(['backend_rows', nrows])
    ws.append(['price_rows', nrows])
    ws.append(['metric'] + [name for name, _ in WINDOWS])
    ws.append([
        'pass_rows',
        f'=COUNTIF(F{COMPARE_START_ROW}:F{compare_end_row},"PASS")',
        f'=COUNTIF(J{COMPARE_START_ROW}:J{compare_end_row},"PASS")',
        f'=COUNTIF(N{COMPARE_START_ROW}:N{compare_end_row},"PASS")',
        f'=COUNTIF(R{COMPARE_START_ROW}:R{compare_end_row},"PASS")',
    ])
    ws.append([
        'empty_rows',
        f'=COUNTIF(F{COMPARE_START_ROW}:F{compare_end_row},"EMPTY")',
        f'=COUNTIF(J{COMPARE_START_ROW}:J{compare_end_row},"EMPTY")',
        f'=COUNTIF(N{COMPARE_START_ROW}:N{compare_end_row},"EMPTY")',
        f'=COUNTIF(R{COMPARE_START_ROW}:R{compare_end_row},"EMPTY")',
    ])
    ws.append([
        'missing_rows',
        f'=COUNTIF(F{COMPARE_START_ROW}:F{compare_end_row},"MISSING")',
        f'=COUNTIF(J{COMPARE_START_ROW}:J{compare_end_row},"MISSING")',
        f'=COUNTIF(N{COMPARE_START_ROW}:N{compare_end_row},"MISSING")',
        f'=COUNTIF(R{COMPARE_START_ROW}:R{compare_end_row},"MISSING")',
    ])
    ws.append([
        'fail_rows',
        f'=COUNTIF(F{COMPARE_START_ROW}:F{compare_end_row},"FAIL")',
        f'=COUNTIF(J{COMPARE_START_ROW}:J{compare_end_row},"FAIL")',
        f'=COUNTIF(N{COMPARE_START_ROW}:N{compare_end_row},"FAIL")',
        f'=COUNTIF(R{COMPARE_START_ROW}:R{compare_end_row},"FAIL")',
    ])
    ws.append([
        'max_abs_diff',
        f'=IFERROR(MAX(E{COMPARE_START_ROW}:E{compare_end_row}),"")',
        f'=IFERROR(MAX(I{COMPARE_START_ROW}:I{compare_end_row}),"")',
        f'=IFERROR(MAX(M{COMPARE_START_ROW}:M{compare_end_row}),"")',
        f'=IFERROR(MAX(Q{COMPARE_START_ROW}:Q{compare_end_row}),"")',
    ])
    ws.append(['overall', '=IF(SUM(B8:E9)=0,"PASS","FAIL")'])
    ws.append([])

    headers = ['row', 'time']
    for suffix in ('10min', '5bar', '1day', '1d30min'):
        headers.extend([
            f'backend_{suffix}',
            f'excel_{suffix}',
            f'diff_{suffix}',
            f'ok_{suffix}',
        ])
    _write_header(ws, headers)

    formulas = _compare_spill_formulas(nrows)
    formatted: list[Any] = []
    for i, formula in enumerate(formulas, start=1):
        if i == 1:
            formatted.append(_styled_cell(ws, formula, number_format=ROW_FMT))
        elif i == 2:
            formatted.append(_styled_cell(ws, formula, number_format=DT_FMT))
        elif i in (6, 10, 14, 18):
            formatted.append(formula)
        else:
            formatted.append(_styled_cell(ws, formula, number_format=NUM_FMT))
    ws.append(formatted)


def process_product(prod_code: str) -> Path:
    prod_path = TEST_2A_DIR / f'{prod_code}.xlsx'
    data = _read_test2a_main(prod_path)
    if data is None:
        raise RuntimeError(f'No cached data for {prod_code} in test_2a')
    trading_days, times, opens, adjs = data
    nrows = len(times)

    product = CNFutures(_product_name(prod_code))
    window_bars = _window_bars(product)

    TEST_4B_DIR.mkdir(parents=True, exist_ok=True)
    out_path = TEST_4B_DIR / f'{prod_code}.xlsx'

    print(f'  {prod_code} ({product.name}): {nrows} rows, rolling windows={window_bars}', flush=True)
    wb = Workbook(write_only=True)

    price_ws = wb.create_sheet(PRICE_SHEET)
    _write_price_sheet(price_ws, trading_days, times, opens, adjs, nrows, window_bars)

    backend_ws = wb.create_sheet(BACKEND_SHEET)
    _write_backend_sheet(backend_ws, product, times)

    compare_ws = wb.create_sheet(COMPARE_SHEET)
    _write_compare_sheet(compare_ws, nrows)

    wb.save(out_path)
    _style_workbook(out_path)
    _patch_dynamic_arrays(out_path, nrows)
    return out_path


def _style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    try:
        for ws in wb.worksheets:
            if ws.max_column > 1:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
            remark_cell = ws.cell(1, 1)
            remark_cell.fill = REMARK_FILL
            remark_cell.font = COMMENT_FONT
            remark_cell.alignment = Alignment(wrap_text=True, vertical='top')
            ws.column_dimensions['A'].width = 22
            for c in 'BCDEFGHIJKLMNOPQR':
                ws.column_dimensions[c].width = 18

        price_ws = wb[PRICE_SHEET]
        price_ws.freeze_panes = 'A3'
        price_formats = {
            'A': DATE_FMT,
            'B': DT_FMT,
            'C': NUM_FMT,
            'D': NUM_FMT,
            'E': NUM_FMT,
            'F': NUM_FMT,
            'G': NUM_FMT,
            'H': NUM_FMT,
            'I': NUM_FMT,
        }
        for col, fmt in price_formats.items():
            price_ws.column_dimensions[col].number_format = fmt

        backend_ws = wb[BACKEND_SHEET]
        backend_ws.freeze_panes = 'A3'
        backend_ws.column_dimensions['A'].number_format = DT_FMT
        for col in 'BCDE':
            backend_ws.column_dimensions[col].number_format = NUM_FMT

        compare_ws = wb[COMPARE_SHEET]
        compare_ws.freeze_panes = f'A{COMPARE_START_ROW}'
        compare_formats = {
            'A': ROW_FMT,
            'B': DT_FMT,
            'C': NUM_FMT,
            'D': NUM_FMT,
            'E': NUM_FMT,
            'G': NUM_FMT,
            'H': NUM_FMT,
            'I': NUM_FMT,
            'K': NUM_FMT,
            'L': NUM_FMT,
            'M': NUM_FMT,
            'O': NUM_FMT,
            'P': NUM_FMT,
            'Q': NUM_FMT,
        }
        for col, fmt in compare_formats.items():
            compare_ws.column_dimensions[col].number_format = fmt
    finally:
        wb.save(path)
        wb.close()


def _sheet_xml_target(zf: zipfile.ZipFile, sheet_name: str) -> str:
    wb_xml = ET.parse(zf.open('xl/workbook.xml'))
    rels_xml = ET.parse(zf.open('xl/_rels/workbook.xml.rels'))
    rel_targets = {rel.get('Id'): rel.get('Target') for rel in rels_xml.getroot()}
    for sh in wb_xml.getroot().findall(f'{{{NS_SHEET}}}sheets/{{{NS_SHEET}}}sheet'):
        if sh.get('name') != sheet_name:
            continue
        r_id = sh.get(f'{{{NS_R}}}id')
        target = rel_targets.get(r_id)
        if target is not None:
            return target.lstrip('/')
    raise RuntimeError(f'Sheet "{sheet_name}" not found in workbook XML')


def _patch_dynamic_arrays(path: Path, nrows: int) -> None:
    last_price_row = DATA_ROW_OFFSET + nrows
    last_compare_row = COMPARE_START_ROW + nrows - 1
    compare_refs = {
        f'{col}{COMPARE_START_ROW}': f'{col}{COMPARE_START_ROW}:{col}{last_compare_row}'
        for col in 'ABCDEFGHIJKLMNOPQR'
    }
    refs_by_sheet = {
        PRICE_SHEET: {
            'E3': f'E3:E{last_price_row}',
            'F3': f'F3:F{last_price_row}',
            'G3': f'G3:G{last_price_row}',
            'H3': f'H3:H{last_price_row}',
            'I3': f'I3:I{last_price_row}',
        },
        COMPARE_SHEET: compare_refs,
    }

    with zipfile.ZipFile(path, 'r') as zf:
        all_files = {name: zf.read(name) for name in zf.namelist()}
        for sheet_name, formula_refs in refs_by_sheet.items():
            target = _sheet_xml_target(zf, sheet_name)
            sheet_root = ET.fromstring(all_files[target])
            wanted_rows = {
                ''.join(ch for ch in cell_ref if ch.isdigit())
                for cell_ref in formula_refs
            }
            for row in sheet_root.findall(f'{{{NS_SHEET}}}sheetData/{{{NS_SHEET}}}row'):
                if row.get('r') not in wanted_rows:
                    continue
                for cell in row.findall(f'{{{NS_SHEET}}}c'):
                    cell_ref = cell.get('r')
                    if cell_ref not in formula_refs:
                        continue
                    formula = cell.find(f'{{{NS_SHEET}}}f')
                    if formula is None:
                        raise RuntimeError(f'{sheet_name}!{cell_ref} formula missing for dynamic-array patch')
                    formula.set('ca', '1')
                    formula.set('t', 'array')
                    formula.set('ref', formula_refs[cell_ref])
                    cell.set('cm', '1')
            all_files[target] = ET.tostring(sheet_root, xml_declaration=True, encoding='UTF-8')

    tmp_fd, tmp_path = tempfile.mkstemp(suffix='.xlsx')
    os.close(tmp_fd)
    try:
        with zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            for name, content in all_files.items():
                zout.writestr(name, content)
        os.replace(tmp_path, path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def _formula_attrs(path: Path, sheet_name: str, cell_ref: str) -> dict[str, str]:
    with zipfile.ZipFile(path, 'r') as zf:
        target = _sheet_xml_target(zf, sheet_name)
        sheet_root = ET.parse(zf.open(target)).getroot()
        for cell in sheet_root.findall(f'.//{{{NS_SHEET}}}c'):
            if cell.get('r') != cell_ref:
                continue
            formula = cell.find(f'{{{NS_SHEET}}}f')
            return {} if formula is None else dict(formula.attrib)
    return {}


def _formula_text(value: Any) -> str:
    text = getattr(value, 'text', value)
    return text if isinstance(text, str) else ''


def main() -> None:
    parser = argparse.ArgumentParser(description='test_4b: verify rolling_mean')
    parser.add_argument('product_codes', nargs='*', default=['A'],
                        help='Product codes (default: A)')
    args = parser.parse_args()

    for code in args.product_codes:
        try:
            process_product(code)
            print(f'  {code}: OK')
        except Exception as e:
            print(f'  {code}: ERROR - {e}')
            import traceback
            traceback.print_exc()
            sys.exit(1)


def test_rolling_mean_workbook_for_a_product():
    out_path = process_product('A')
    wb = load_workbook(out_path, read_only=False, data_only=False)
    try:
        assert wb.sheetnames == [PRICE_SHEET, BACKEND_SHEET, COMPARE_SHEET]
        price_ws = wb[PRICE_SHEET]
        backend_ws = wb[BACKEND_SHEET]
        compare_ws = wb[COMPARE_SHEET]
        assert [str(r) for r in price_ws.merged_cells.ranges] == ['A1:I1']
        assert [str(r) for r in backend_ws.merged_cells.ranges] == ['A1:E1']
        assert [str(r) for r in compare_ws.merged_cells.ranges] == ['A1:R1']
        for ws in (price_ws, backend_ws, compare_ws):
            assert ws['A1'].font.italic is True
            assert ws['A1'].font.bold is False
        assert price_ws.cell(2, 1).value == 'trading_day'
        assert price_ws.cell(2, 6).value == 'ROLLING_MEAN_10MIN'
        assert price_ws.cell(2, 9).value == 'ROLLING_MEAN_1DAY_30MIN'
        assert backend_ws.cell(2, 5).value == 'ROLLING_MEAN_1DAY_30MIN'
        assert compare_ws.cell(2, 1).value == 'TOLERANCE'
        assert compare_ws.cell(5, 1).value == 'metric'
        assert compare_ws.cell(11, 2).value == '=IF(SUM(B8:E9)=0,"PASS","FAIL")'
        rolling_formula = _formula_text(price_ws.cell(BACKEND_DATA_START_ROW, 6).value)
        assert 'MAP(SEQUENCE(ROWS(x))' in rolling_formula
        assert 'minp,5' in rolling_formula
        rolling_1d30_formula = _formula_text(price_ws.cell(BACKEND_DATA_START_ROW, 9).value)
        assert 'w,375' in rolling_1d30_formula
        assert 'minp,187' in rolling_1d30_formula
        backend_10min_formula = _formula_text(compare_ws.cell(COMPARE_START_ROW, 3).value)
        assert 'BACKEND!B3:B' in backend_10min_formula
        assert 'IF(LEN(r&"")=0,"",r)' in backend_10min_formula
        excel_1d30_formula = _formula_text(compare_ws.cell(COMPARE_START_ROW, 16).value)
        assert 'PRICE!I3:I' in excel_1d30_formula
        assert price_ws.cell(BACKEND_DATA_START_ROW, 1).number_format == DATE_FMT
        assert price_ws.cell(BACKEND_DATA_START_ROW, 6).number_format == NUM_FMT
        assert compare_ws.cell(COMPARE_START_ROW, 2).number_format == DT_FMT
        assert compare_ws.cell(COMPARE_START_ROW, 3).number_format == NUM_FMT
        price_last_row = DATA_ROW_OFFSET + (price_ws.max_row - DATA_ROW_OFFSET)
        assert _formula_attrs(out_path, PRICE_SHEET, 'E3') == {
            'ca': '1', 't': 'array', 'ref': f'E3:E{price_last_row}'
        }
        assert _formula_attrs(out_path, PRICE_SHEET, 'I3') == {
            'ca': '1', 't': 'array', 'ref': f'I3:I{price_last_row}'
        }
        compare_last_row = COMPARE_START_ROW + (price_ws.max_row - DATA_ROW_OFFSET) - 1
        assert _formula_attrs(out_path, COMPARE_SHEET, f'A{COMPARE_START_ROW}') == {
            'ca': '1', 't': 'array', 'ref': f'A{COMPARE_START_ROW}:A{compare_last_row}'
        }
        assert _formula_attrs(out_path, COMPARE_SHEET, f'R{COMPARE_START_ROW}') == {
            'ca': '1', 't': 'array', 'ref': f'R{COMPARE_START_ROW}:R{compare_last_row}'
        }
    finally:
        wb.close()


if __name__ == '__main__':
    main()
