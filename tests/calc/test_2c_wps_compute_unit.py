from pathlib import Path

from openpyxl import Workbook, load_workbook

from tests.calc import test_2c_wps_compute as wps_compute


def _make_test2a_workbook(path: Path) -> None:
    wb = Workbook()
    main = wb.active
    main.title = 'MAIN'
    main.cell(2, 1, 'trade_time')
    main.cell(2, 24, 'adjustment_mul')
    main.cell(3, 24, 1.0)

    sw = wb.create_sheet('_SWITCHES')
    sw.cell(2, 1, 'trading_day')
    sw.cell(3, 1, '2026-01-01')
    sw.cell(4, 1, '2026-01-02')
    wb.save(path)
    wb.close()


def test_patch_flag_writes_pre_open_save_handshake_formulas(tmp_path):
    path = tmp_path / 'A.xlsx'
    _make_test2a_workbook(path)

    wps_compute.patch_flag(path, save_target=7)

    wb = load_workbook(path, data_only=False)
    sw = wb['_SWITCHES']
    assert sw['Z1'].value == '=IF(ISNUMBER(MAIN!X4),"DONE","CALC")'
    assert sw['Z2'].value == '=IF(Z1="DONE",7,0)'
    wb.close()


def test_wps_save_and_check_requires_saved_target_and_done(monkeypatch, tmp_path):
    path = tmp_path / 'A.xlsx'
    calls = {'n': 0}

    def fake_run(*args, **kwargs):
        return None

    def fake_z2(_path):
        calls['n'] += 1
        return 3 if calls['n'] >= 2 else 0

    monkeypatch.setattr(wps_compute.subprocess, 'run', fake_run)
    monkeypatch.setattr(wps_compute, 'read_z2', fake_z2)
    monkeypatch.setattr(wps_compute, 'read_z1', lambda _path: 'DONE')
    monkeypatch.setattr(wps_compute.time, 'sleep', lambda _seconds: None)

    assert wps_compute.wps_save_and_check(path, save_target=3, timeout=1) == (True, True)


def test_next_save_target_is_greater_than_cached_count(monkeypatch, tmp_path):
    path = tmp_path / 'A.xlsx'

    monkeypatch.setattr(wps_compute, 'read_z2', lambda _path: 10)
    monkeypatch.setattr(wps_compute.time, 'time', lambda: 5)

    assert wps_compute.next_save_target(path) == 11


def test_wps_save_and_check_rejects_saved_count_without_done(monkeypatch, tmp_path):
    path = tmp_path / 'A.xlsx'
    times = iter([0, 0.5, 1.5])

    monkeypatch.setattr(wps_compute.subprocess, 'run', lambda *args, **kwargs: None)
    monkeypatch.setattr(wps_compute, 'read_z2', lambda _path: 3)
    monkeypatch.setattr(wps_compute, 'read_z1', lambda _path: 'CALC')
    monkeypatch.setattr(wps_compute.time, 'sleep', lambda _seconds: None)
    monkeypatch.setattr(wps_compute.time, 'time', lambda: next(times))

    assert wps_compute.wps_save_and_check(path, save_target=3, timeout=1) == (True, False)
